#!/usr/bin/env python3
"""Dev-time generator for the Ticket 02 axis-neutral routing fixtures (R & C).

⛔ 测试运行时绝不 import 本脚本：与 generate_prepare_evidence_fixture.py /
generate_task_orchestration_e2e.py 同策略 —— 预生成合成工作簿提交到
tests/_fixtures/routing/，运行时只复制 checked-in 副本。

为什么 openpyxl 之外还做一步 shared-string 归一化：
  officecli 1.0.145 对 openpyxl 生成的 inlineStr 单元格报告
  format.type == "InlineString"，而 detect_header_rows 只按 "SharedString"
  计数（Case 010 契约锁定，Ticket 01 已实测并记录该机械事实）→ 不归一化时
  表头带检测永不触发，Layer 2 的「表头带 + 数据起始行」证据断言会失败。
  本生成器先用 openpyxl 确定性构建逻辑网格（值、真实 merged ranges、边框、
  列宽），再用 zip 级机械后处理把全部字符串格转为共享字符串（t="s"）并补
  sharedStrings 部件与关系 —— 数值格原样保留。两步都是确定性的；工作簿的
  每个值都由本文件单一定义。

Fixture 设计（data-neutral：无 MXP/ATLAS/TCL/Egypt/Algeria/真实型号/真实报价/
真实客户名/中国业务词；产品占位词 12K/18K/24K 为 spec 认可的通用 token）：

  Fixture R — Grouped Row Grid（row-record；报价形态，供 obvious_grid/rows）：
    Row 1:   A1:F1 合并标题 "Product Quotation List"
    Row 2:   六列表头 (Type|Model|Capacity|Qty|Unit Price|Panel looking)
    Rows 3-8: 两组重复记录行（每组 3 条）；A、F 纵向组 merge（真实 merged
              ranges，锚点值在组首行）；Unit Price 列记录行留空（与 Case 010
              ATLAS 数据行密度一致：记录行 5 非空 < 表头带 6 → band=[2]/3）
    Rows 9-11: Total / Features / Notes 外围行
    期望: grid_record / record_axis=rows / fillspec / ["obvious_grid"]

  Fixture C — Column Record Matrix（column-record；参数表形态，
    供 obvious_grid/columns）：
    Row 1:   A1:D1 合并标题 "Product Parameter Matrix"（真实 merged range）
    Row 2:   表头示意行 (Parameter | 12K | 18K | 24K)
    Rows 3-8: 字段标签列(A) × 产品列(B/C/D) 数据行（首行即数值多数的数据行
              → band=[2]/3；标题行稀疏 → 不吸入）
    期望: grid_record / record_axis=columns / fillspec / ["obvious_grid"]

用法:
  python tests/_fixtures/routing/generate_fixtures.py
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape, unescape

import openpyxl
from openpyxl.styles import Border, Side

OUT_DIR = Path(__file__).resolve().parent
FIX_R = OUT_DIR / "fixture_r_grouped_row_grid.xlsx"
FIX_C = OUT_DIR / "fixture_c_column_record_matrix.xlsx"

_SHEET_NAME = "Sheet1"

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _thin_borders(ws, cell_range: str) -> None:
    """Apply a real thin border to every cell in `cell_range` (A1:F11 form).

    边框是真实样式 → compute_bbox_facts 的 style_bbox 出现（evidence 用
    `style_bbox` if present 分支）。"""
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = _BORDER


def _convert_inline_strings(xml: str, sid) -> str:
    """把 sheet XML 中全部 inlineStr 字符串格改写为共享字符串引用。

    只动 字符串格：`<c r="A1" t="inlineStr" s="1"><is><t>..</t></is></c>`
    → `<c r="A1" s="1" t="s"><v>IDX</v></c>`；数值格原样保留。属性顺序
    无关（分别取 r/s/t 与 <t> 文本）。"""
    cell_re = re.compile(r"<c\b[^>]*>.*?</c>|<c\b[^>]*/>", re.S)

    def repl(m):
        tag = m.group(0)
        if "inlineStr" not in tag:
            return tag
        rm = re.search(r'\br="([A-Z]+\d+)"', tag)
        sm = re.search(r'\bs="(\d+)"', tag)
        tm = re.search(r"<is>(?:<r>)?<t(?:\s[^>]*)?>(.*?)</t>", tag, re.S)
        if not (rm and tm):
            return tag
        text = unescape(tm.group(1))
        s = f' s="{sm.group(1)}"' if sm else ""
        return f'<c r="{rm.group(1)}"{s} t="s"><v>{sid(text)}</v></c>'

    return cell_re.sub(repl, xml)


def _shared_stringify(src: Path, dst: Path) -> None:
    """zip 级共享字符串归一化（机械后处理，确定性）。

    officecli 1.0.145 把 openpyxl 的 inlineStr 报为 format.type==
    "InlineString"，detect_header_rows 只数 "SharedString"（Case 010 契约，
    Ticket 01 实测记录）→ 表头带在 openpyxl 原生输出上永不触发。本步把全部
    字符串格转为 t="s" 共享字符串并补 sharedStrings 部件 + content-type +
    workbook 关系，使 header band / 候选列头在真实 prepare 全链路中可达。"""
    tokens: list[str] = []
    index: dict[str, int] = {}

    def sid(text: str) -> int:
        if text not in index:
            index[text] = len(tokens)
            tokens.append(text)
        return index[text]

    with zipfile.ZipFile(src) as zin:
        names = zin.namelist()
        out: dict[str, bytes] = {}
        for name in names:
            data = zin.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = data.decode("utf-8")
                data = _convert_inline_strings(xml, sid).encode("utf-8")
            out[name] = data

        ct = out["[Content_Types].xml"].decode("utf-8")
        if "sharedStrings.xml" not in ct:
            ct = ct.replace(
                "</Types>",
                '<Override PartName="/xl/sharedStrings.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sharedStrings+xml"/></Types>')
            out["[Content_Types].xml"] = ct.encode("utf-8")

        rels_name = "xl/_rels/workbook.xml.rels"
        rels = out[rels_name].decode("utf-8")
        if "sharedStrings.xml" not in rels:
            ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
            rid = (max(ids) + 1) if ids else 1
            rels = rels.replace(
                "</Relationships>",
                f'<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats'
                '.org/officeDocument/2006/relationships/sharedStrings" '
                'Target="sharedStrings.xml"/></Relationships>')
            out[rels_name] = rels.encode("utf-8")

        si = "".join(f"<si><t>{escape(t)}</t></si>" for t in tokens)
        sst = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{len(tokens)}" uniqueCount="{len(tokens)}">\n{si}\n</sst>'
        )
        out["xl/sharedStrings.xml"] = sst.encode("utf-8")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out.items():
            z.writestr(name, data)


def _build_fixture_r() -> None:
    """Fixture R — Grouped Row Grid（row-record）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _SHEET_NAME

    ws.merge_cells("A1:F1")
    ws["A1"] = "Product Quotation List"

    for col, h in zip("ABCDEF",
                      ("Type", "Model", "Capacity", "Qty", "Unit Price",
                       "Panel looking")):
        ws[f"{col}2"] = h

    # 两组重复记录行；A、F 纵向组 merge（真实 merged ranges，锚点在组首行）
    groups = [
        ("Wall", "Classic", 3, [("Model L1", "9000Btu", 2),
                                ("Model L2", "12000Btu", 3),
                                ("Model L3", "18000Btu", 4)]),
        ("Portable", "Modern", 6, [("Model M1", "7000Btu", 1),
                                   ("Model M2", "9000Btu", 2),
                                   ("Model M3", "12000Btu", 2)]),
    ]
    for gtype, panel, start, recs in groups:
        end = start + len(recs) - 1
        ws.merge_cells(f"A{start}:A{end}")
        ws[f"A{start}"] = gtype
        ws.merge_cells(f"F{start}:F{end}")
        ws[f"F{start}"] = panel
        for i, (model, capacity, qty) in enumerate(recs):
            r = start + i
            ws[f"B{r}"] = model
            ws[f"C{r}"] = capacity
            ws[f"D{r}"] = qty
            # E (Unit Price) 记录行留空 —— 与 Case 010 ATLAS 数据行密度一致

    # 外围结构：Total / Features / Notes（不改变 Grid 身份）
    ws["A9"] = "Total"
    ws.merge_cells("B9:F9")
    ws["B9"] = "Sum of listed prices"
    ws["A10"] = "Features"
    ws.merge_cells("B10:F10")
    ws["B10"] = "Grouped by product family"
    ws["A11"] = "Notes"
    ws.merge_cells("B11:F11")
    ws["B11"] = "Placeholder note only"

    _thin_borders(ws, "A1:F11")
    for col, w in zip("ABCDEF", (10, 12, 12, 8, 12, 14)):
        ws.column_dimensions[col].width = w

    wb.save(FIX_R)
    _shared_stringify(FIX_R, FIX_R)
    print(f"wrote {FIX_R}")


def _build_fixture_c() -> None:
    """Fixture C — Column Record Matrix（column-record)。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _SHEET_NAME

    # Row1: 合并标题（稀疏行，表头带跳过；表格 [1] 行也是 CSV 首行，避免
    # utc-8-sig BOM 落到表头示意行上）
    ws.merge_cells("A1:D1")
    ws["A1"] = "Product Parameter Matrix"

    # Row2 表头示意行：字段标签列名 + 产品 record 列（12K/18K/24K 通用占位词）
    for col, h in zip("ABCD", ("Parameter", "12K", "18K", "24K")):
        ws[f"{col}2"] = h

    for i, (label, b, c, d) in enumerate(
            [("Capacity", 12000, 18000, 24000),
             ("EER", 11.0, 10.8, 11.6),
             ("Sound", 42, 45, 48),
             ("Type", "Wall", "Wall", "Wall"),
             ("Compressor", "VS1", "VS2", "VS3"),
             ("Refrigerant", "R32", "R32", "R32")],
            start=3):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = b
        ws[f"C{i}"] = c
        ws[f"D{i}"] = d

    _thin_borders(ws, "A1:D8")
    for col, w in zip("ABCD", (14, 10, 10, 10)):
        ws.column_dimensions[col].width = w

    wb.save(FIX_C)
    _shared_stringify(FIX_C, FIX_C)
    print(f"wrote {FIX_C}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    _build_fixture_r()
    _build_fixture_c()


if __name__ == "__main__":
    main()