#!/usr/bin/env python3
"""Dev-time generator for the Ticket 02 (P0-02) hierarchical matrix canonical
fixtures — `source_parameter_book.xlsx` (internal-facing, WITH extra
internal-only rows) and `target_template.xlsx` (customer-facing template,
WITHOUT the internal rows).

⛔ 测试运行时绝不 import 本脚本：与 routing/generate_fixtures.py 同策略 ——
预生成合成工作簿提交到 tests/_fixtures/hierarchical_matrix/，运行时只复制
checked-in 副本，绝不在测试里现场生成。

为什么 openpyxl 之外还做一步 shared-string 归一化（与 routing 生成器同一
两步模式，逐字复用其机械后处理）：
  officecli 1.0.145 对 openpyxl 生成的 inlineStr 单元格报告
  format.type == "InlineString"，而 detect_header_rows 只按 "SharedString"
  计数 → 不归一化时表头带永不触发。本生成器先用 openpyxl 确定性构建逻辑
  网格（值、真实 merged ranges、边框、列宽），再用 zip 级机械后处理把全部
  字符串格转为共享字符串（t="s"）并补 sharedStrings 部件与关系 —— 数值格
  原样保留。两步都是确定性的；工作簿的每个值都由本文件单一定义。

Fixture 结构契约（spec D9，data-neutral；禁 Egypt/TCL/MXP/ATLAS/真实 Z 码/
真实客户名/真实型号 —— 只复刻结构复杂度）：

  源书 sheet "Params"（内部面，多 internal-only 行）:
    Row 1:  A1:F1 合并标题 "Product Parameter Book (internal)"
    Row 2:  表头示意行 (Section | Parameter | Unit | 12K | 18K | 24K)
    Row 3-6: Cooling section — A3:A6 真实纵向合并 "Cooling"：
             Capacity/Btu/h (12000/18000/24000)、Capacity/W (3500/5300/7000)、
             EER/Btu/Wh (11/10.8/11.6)、Operation/Mode
             ("Heating pump" ← 受控翻译源值 → 目标 "Cooling and Heating")
    Row 7-9: Heating section — A7:A9 真实纵向合并 "Heating"：
             Capacity/Btu/h (14000/21000/28000)、COP/W/W (3.6/3.8/4.0)、
             Sound/dB(A) (42/45/48)
    Row 10:  Voltage | Supply | V/Ph/Hz | " 220 "（首尾空白 trim 用例）
    Row 11:  Internal | Cost Code | USD | "内部成本" ×3 —— internal-only 行，
             目标模板没有，field_map 不映射 → 永不进入最终 workbook

  目标模板 sheet "Template"（客户面，无 internal 行）:
    Row 1:  A1:F1 合并标题 "Product Parameter Matrix (template)"
    Row 2:  表头示意行（同源）
    Row 3:  A3:F3 合并 note 行 "Data to be filled by application engineer"
            （目标比源多一行 → 源/目标行号刻意错位，locator 不依赖绝对行序）
    Row 4-7: Cooling section — A4:A7 合并 "Cooling"，同标签行（D/E/F 空）
    Row 8-10: Heating section — A8:A10 合并 "Heating"，同标签行（D/E/F 空）
    Row 11: Voltage | Supply | V/Ph/Hz（D/E/F 空）
    Row 12:  footer "Note: values are data-neutral and generic."

  层级身份 = (Section=A, Parameter=B, Unit=C) 三元组；B 重复 (Capacity ×3)、
  C 最终消歧 (Btu/h vs W)；flatten (STANDARD 纵向合并传播) 把合并锚点
  "Cooling"/"Heating" 传播进每个子行的 A 列 → composite match 键 A 生效。

  ⚠️ 受控翻译用英文源值 "Heating pump"（spec 许可的通用示例
  "Heat Pump" → "Cooling and Heating" 变体）：officecli 1.0.145 的 get/batch
  对非 ASCII 共享字符串值返回/写入 XML 数字字符引用（实测
  `text: "&#23485;&#29255;"`），compile 从展平 CSV 物化出的正是转义字面量，
  controlled_translation 精确匹配无法命中；若强制中文受控值，最终 workbook
  会携带转义串（非真实 宽片）。因此中文数据点放在不映射的 internal-only 行
  （"内部成本"），受控翻译走 ASCII 全链（编译 → 执行 → readback → 语义扫描
  全部实测一致）。这不是削弱 — ticket 的 "or" 使 spec 许可的英文示例完全合规。

用法:
  python tests/_fixtures/hierarchical_matrix/generate_fixtures.py
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape, unescape

import openpyxl
from openpyxl.styles import Border, Side

OUT_DIR = Path(__file__).resolve().parent
SRC_XLSX = OUT_DIR / "source_parameter_book.xlsx"
TGT_XLSX = OUT_DIR / "target_template.xlsx"

SRC_SHEET = "Params"
TGT_SHEET = "Template"

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _thin_borders(ws, cell_range: str) -> None:
    """Apply a real thin border to every cell in `cell_range` (A1:F12 form)."""
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = _BORDER


def _convert_inline_strings(xml: str, sid) -> str:
    """把 sheet XML 中全部 inlineStr 字符串格改写为共享字符串引用。

    只动字符串格：`<c r="A1" t="inlineStr" s="1"><is><t>..</t></is></c>`
    → `<c r="A1" s="1" t="s"><v>IDX</v></c>`；数值格原样保留。属性顺序
    无关（分别取 r/s/t 与 <t> 文本）。"""
    cell_re = re.compile(r"<c\b[^>]*>.*?</c>|<c\b[^>]*/>", re.S)

    def repl(m):
        tag = m.group(0)
        if "inlineStr" not in tag:
            return tag
        rm = re.search(r'\br="([A-Z]+\d+)"', tag)
        sm = re.search(r'\bs="(\d+)"', tag)
        tm = re.search(r"<is>(?:<r>)?<t(?:\s[^>]*)?>(.*?)</t>", tag, re.S)
        if not (rm and tm):
            return tag
        text = unescape(tm.group(1))
        s = f' s="{sm.group(1)}"' if sm else ""
        return f'<c r="{rm.group(1)}"{s} t="s"><v>{sid(text)}</v></c>'

    return cell_re.sub(repl, xml)


def _shared_stringify(src: Path, dst: Path) -> None:
    """zip 级共享字符串归一化（机械后处理，确定性）—— 与
    routing/generate_fixtures.py 的两步模式完全一致。"""
    tokens: list[str] = []
    index: dict[str, int] = {}

    def sid(text: str) -> int:
        if text not in index:
            index[text] = len(tokens)
            tokens.append(text)
        return index[text]

    with zipfile.ZipFile(src) as zin:
        names = zin.namelist()
        out: dict[str, bytes] = {}
        for name in names:
            data = zin.read(name)
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = data.decode("utf-8")
                data = _convert_inline_strings(xml, sid).encode("utf-8")
            out[name] = data

        ct = out["[Content_Types].xml"].decode("utf-8")
        if "sharedStrings.xml" not in ct:
            ct = ct.replace(
                "</Types>",
                '<Override PartName="/xl/sharedStrings.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sharedStrings+xml"/></Types>')
            out["[Content_Types].xml"] = ct.encode("utf-8")

        rels_name = "xl/_rels/workbook.xml.rels"
        rels = out[rels_name].decode("utf-8")
        if "sharedStrings.xml" not in rels:
            ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels)]
            rid = (max(ids) + 1) if ids else 1
            rels = rels.replace(
                "</Relationships>",
                f'<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats'
                '.org/officeDocument/2006/relationships/sharedStrings" '
                'Target="sharedStrings.xml"/></Relationships>')
            out[rels_name] = rels.encode("utf-8")

        si = "".join(f"<si><t>{escape(t)}</t></si>" for t in tokens)
        sst = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{len(tokens)}" uniqueCount="{len(tokens)}">\n{si}\n</sst>'
        )
        out["xl/sharedStrings.xml"] = sst.encode("utf-8")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in out.items():
            z.writestr(name, data)


# 层级参数表行（(B=Parameter, C=Unit, D/E/F=12K/18K/24K) 值元组），
# 每条记录归属 (section, 起点行, 合并跨度)。
SRC_SECTIONS = [
    ("Cooling", 3, ["Capacity", "Btu/h", 12000, 18000, 24000],
                 ["Capacity", "W",     3500,  5300,  7000],
                 ["EER",      "Btu/Wh", 11,   10.8,  11.6],
                 ["Operation", "Mode", "Heating pump", "Heating pump",
                  "Heating pump"]),   # 受控翻译源值（spec 许可的英文示例）
    ("Heating", 7, ["Capacity", "Btu/h", 14000, 21000, 28000],
                  ["COP",      "W/W",   3.6,   3.8,   4.0],
                  ["Sound",    "dB(A)", 42,    45,    48]),
]
SRC_SOLO_ROWS = [
    # (B, C, D, E, F)
    ("Supply", "V/Ph/Hz", " 220 ", " 220 ", " 220 "),   # trim 用例（首尾空白）
]
SRC_INTERNAL_ROW = ["Cost Code", "USD",
                    "内部成本", "内部成本", "内部成本"]  # internal-only（中文内部值）

TGT_SECTIONS = [
    ("Cooling", 4, ["Capacity", "Btu/h"], ["Capacity", "W"],
                   ["EER", "Btu/Wh"], ["Operation", "Mode"]),
    ("Heating", 8, ["Capacity", "Btu/h"], ["COP", "W/W"], ["Sound", "dB(A)"]),
]
TGT_SOLO_ROWS = [("Supply", "V/Ph/Hz")]
TGT_FOOTER = "Note: values are data-neutral and generic."


def _header_row(ws, row: int) -> None:
    for col, h in zip("ABCDEF",
                      ("Section", "Parameter", "Unit", "12K", "18K", "24K")):
        ws[f"{col}{row}"] = h


def _build_source() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SRC_SHEET

    ws.merge_cells("A1:F1")
    ws["A1"] = "Product Parameter Book (internal)"
    _header_row(ws, 2)

    for section, start, *rows in SRC_SECTIONS:
        end = start + len(rows) - 1
        ws.merge_cells(f"A{start}:A{end}")
        ws[f"A{start}"] = section
        for i, (b, c, d, e, f_) in enumerate(rows):
            r = start + i
            ws[f"B{r}"] = b
            ws[f"C{r}"] = c
            ws[f"D{r}"] = d
            ws[f"E{r}"] = e
            ws[f"F{r}"] = f_

    solo_start = 10
    for i, (b, c, d, e, f_) in enumerate(SRC_SOLO_ROWS):
        r = solo_start + i
        ws[f"A{r}"] = "Voltage"
        ws[f"B{r}"] = b
        ws[f"C{r}"] = c
        ws[f"D{r}"] = d
        ws[f"E{r}"] = e
        ws[f"F{r}"] = f_

    int_row = solo_start + len(SRC_SOLO_ROWS)
    ws["A" + str(int_row)] = "Internal"
    for col, v in zip("BCDEF", SRC_INTERNAL_ROW):
        ws[f"{col}{int_row}"] = v

    _thin_borders(ws, "A1:F11")
    for col, w in zip("ABCDEF", (12, 12, 12, 10, 10, 10)):
        ws.column_dimensions[col].width = w

    wb.save(SRC_XLSX)
    _shared_stringify(SRC_XLSX, SRC_XLSX)
    print(f"wrote {SRC_XLSX}")


def _build_target() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TGT_SHEET

    ws.merge_cells("A1:F1")
    ws["A1"] = "Product Parameter Matrix (template)"
    _header_row(ws, 2)

    ws.merge_cells("A3:F3")
    ws["A3"] = "Data to be filled by application engineer"

    for section, start, *rows in TGT_SECTIONS:
        end = start + len(rows) - 1
        ws.merge_cells(f"A{start}:A{end}")
        ws[f"A{start}"] = section
        for i, (b, c) in enumerate(rows):
            r = start + i
            ws[f"B{r}"] = b
            ws[f"C{r}"] = c
            # D/E/F 留空 —— matrix 填充

    solo_start = 11
    for i, (b, c) in enumerate(TGT_SOLO_ROWS):
        r = solo_start + i
        ws[f"A{r}"] = "Voltage"
        ws[f"B{r}"] = b
        ws[f"C{r}"] = c

    footer_row = solo_start + len(TGT_SOLO_ROWS)
    ws[f"A{footer_row}"] = TGT_FOOTER

    _thin_borders(ws, "A1:F12")
    for col, w in zip("ABCDEF", (12, 12, 12, 10, 10, 10)):
        ws.column_dimensions[col].width = w

    wb.save(TGT_XLSX)
    _shared_stringify(TGT_XLSX, TGT_XLSX)
    print(f"wrote {TGT_XLSX}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    _build_source()
    _build_target()


if __name__ == "__main__":
    main()